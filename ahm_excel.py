# -*- coding: utf-8 -*-
"""Lecture d'un fichier AHM 565 au format Excel (.xlsx).

Principe : chaque feuille Excel est présentée au lecteur PDF existant (ahm565.py) comme une page, avec des mots
positionnés d'après la grille de colonnes et de lignes. Le même code lit donc le PDF et l'Excel : mêmes contrôles,
mêmes avertissements, même structure de résultat. Le module ne contient aucune valeur de compagnie.
"""
import io
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import ahm565 as AHM

PAGE_W = 595.2          # largeur de page A4 du PDF : les colonnes Excel sont mises à cette échelle
X_A, X_B = 1.06, 62.0   # calage affine des abscisses Excel sur celles du PDF (marges de la page imprimée)
ROW_SCALE = 0.6         # hauteur des lignes Excel ramenée à l'interligne d'un PDF


class _Page:
    def __init__(self, words):
        self._words = words          # (x0, y0, x1, y1, texte)

    def get_text(self, mode="text"):
        if mode == "words":
            return list(self._words)
        rows = {}
        for w in self._words:
            rows.setdefault(round(w[1] / 4), []).append(w)
        return "\n".join(" ".join(x[4] for x in sorted(r, key=lambda z: z[0])) for _, r in sorted(rows.items()))


def _sheet_words(ws):
    ncol = ws.max_column
    widths = []
    for c in range(1, ncol + 1):
        d = ws.column_dimensions.get(get_column_letter(c))
        widths.append(d.width if d is not None and d.width else (ws.sheet_format.defaultColWidth or 8.43))
    tot = sum(widths) or 1.0
    xs, acc = [], 0.0
    for w in widths:
        xs.append(((acc / tot * PAGE_W) * X_A + X_B, ((acc + w) / tot * PAGE_W) * X_A + X_B))
        acc += w
    dflt = ws.sheet_format.defaultRowHeight or 15.0
    ys, acc = {}, 0.0
    for r in range(1, ws.max_row + 1):
        h = ws.row_dimensions[r].height or dflt
        ys[r] = acc * ROW_SCALE
        acc += h
    words = []
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
                v = int(v)
            txt = str(v)
            x0, x1 = xs[cell.column - 1]
            for k, line in enumerate(txt.splitlines() or [txt]):
                toks = line.split()
                if not toks:
                    continue
                total = sum(len(t) for t in toks) + len(toks) - 1
                pos = 0
                span = max(x1 - x0, 1.0)
                for t in toks:
                    a = x0 + span * pos / total
                    b = x0 + span * (pos + len(t)) / total if len(toks) > 1 else x1
                    y0 = ys[cell.row] + 4.0 * k
                    words.append((a, y0, b, y0 + 8.0, t))
                    pos += len(t) + 1
    return words


class XlsxDoc(AHM.Doc):
    """Même interface que Doc (PDF), construite à partir d'un classeur Excel."""
    def __init__(self, data: bytes = None, path: str = None):
        wb = load_workbook(io.BytesIO(data) if data is not None else path, data_only=True, read_only=False)
        self.sheet_names = list(wb.sheetnames)
        self.pdf = [_Page(_sheet_words(wb[n])) for n in self.sheet_names]
        self.n = len(self.pdf)
        self.warn = []
        self.pages = {}
        self._lines = {}
        self._text = {}
        for i, pg in enumerate(self.pdf):
            t = pg.get_text()
            self._text[i] = t
            flat = re.sub(r'\s+', ' ', t)
            for key, pat in AHM.MARKERS:
                if re.search(pat, flat, re.I if key not in ('cabin',) else 0):
                    self.pages.setdefault(key, []).append(i)

    def lines(self, i):
        if i not in self._lines:
            self._lines[i] = AHM.page_lines(self.pdf[i])
        return self._lines[i]


def is_ahm565_xlsx(data: bytes) -> bool:
    try:
        d = XlsxDoc(data=data)
        head = ' '.join(d.text(i) for i in range(min(d.n, 12)))
        return 'AHM' in head and ('565' in head) and ('SEMI' in head or 'EDP' in head)
    except Exception:
        return False


def parse_all_xlsx(data: bytes, filename: str = ""):
    return AHM.parse_all(doc=XlsxDoc(data=data), filename=filename)
